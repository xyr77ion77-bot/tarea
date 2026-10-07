#!/usr/bin/env python3
"""
generar_presentacion.py
=======================
Genera una plantilla de presentación en PowerPoint (.pptx) con una estructura
genérica lista para personalizar:

    1. Portada
    2. Índice
    3. Introducción
    4. Objetivos
    5. Desarrollo (tema 1)
    6. Desarrollo (tema 2)
    7. Resultados (con gráfico de barras)
    8. Datos clave (con tabla)
    9. Conclusiones
   10. Cierre / Preguntas

Uso:
    python generar_presentacion.py            # crea "presentacion.pptx"
    python generar_presentacion.py salida.pptx  # nombre personalizado

Requisitos:
    pip install python-pptx
"""

import sys
from datetime import date

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

# ---------------------------------------------------------------------------
# CONFIGURACIÓN - edita estos valores para personalizar tu presentación
# ---------------------------------------------------------------------------

TITULO = "Título de la Presentación"
SUBTITULO = "Subtítulo, evento o materia"
AUTOR = "Autor: Tu Nombre"
FECHA = date.today().strftime("%d/%m/%Y")
CONTACTO = "correo@ejemplo.com"

COLORES = {
    "primario": RGBColor(0x1F, 0x38, 0x64),   # azul oscuro (fondos, títulos)
    "acento": RGBColor(0x2E, 0x9C, 0xCA),     # azul claro (barras, detalles)
    "texto": RGBColor(0x33, 0x33, 0x33),      # gris oscuro (cuerpo)
    "claro": RGBColor(0xF2, 0xF4, 0xF7),      # gris claro (tarjetas)
    "blanco": RGBColor(0xFF, 0xFF, 0xFF),
}

FUENTE = "Calibri"
ANCHO, ALTO = Inches(13.333), Inches(7.5)     # 16:9

# ---------------------------------------------------------------------------
# FUNCIONES AUXILIARES
# ---------------------------------------------------------------------------


def fondo(slide, color):
    """Rellena el fondo de la diapositiva con un color."""
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def rectangulo(slide, left, top, width, height, color, sin_borde=True):
    """Añade un rectángulo de color (usado como barra, tarjeta o banda)."""
    forma = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    forma.fill.solid()
    forma.fill.fore_color.rgb = color
    if sin_borde:
        forma.line.fill.background()
    return forma


def cuadro_texto(slide, left, top, width, height, texto, tamaño=18,
                 negrita=False, color=None, alineacion=PP_ALIGN.LEFT,
                 anchor=MSO_ANCHOR.TOP):
    """Añade un cuadro de texto con una sola línea o párrafo."""
    color = color or COLORES["texto"]
    caja = slide.shapes.add_textbox(left, top, width, height)
    tf = caja.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    párrafo = tf.paragraphs[0]
    párrafo.alignment = alineacion
    run = párrafo.add_run()
    run.text = texto
    run.font.size = Pt(tamaño)
    run.font.bold = negrita
    run.font.color.rgb = color
    run.font.name = FUENTE
    return caja


def lista_viñetas(slide, left, top, width, height, elementos,
                  tamaño=18, color=None):
    """Añade una lista con viñetas ('•' nivel 1, '–' nivel 2)."""
    color = color or COLORES["texto"]
    caja = slide.shapes.add_textbox(left, top, width, height)
    tf = caja.text_frame
    tf.word_wrap = True
    for i, (nivel, texto) in enumerate(elementos):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(10)
        p.level = nivel
        run = p.add_run()
        prefijo = "   " * nivel + ("•  " if nivel == 0 else "–  ")
        run.text = prefijo + texto
        run.font.size = Pt(tamaño - 2 * nivel)
        run.font.color.rgb = color
        run.font.name = FUENTE
    return caja


def encabezado(slide, titulo, numero):
    """Barra de título estándar para diapositivas de contenido + paginación."""
    fondo(slide, COLORES["blanco"])
    rectangulo(slide, 0, 0, ANCHO, Inches(1.1), COLORES["primario"])
    rectangulo(slide, 0, Inches(1.1), ANCHO, Inches(0.08), COLORES["acento"])
    cuadro_texto(slide, Inches(0.6), Inches(0.18), Inches(11), Inches(0.8),
                 titulo, tamaño=30, negrita=True, color=COLORES["blanco"],
                 anchor=MSO_ANCHOR.MIDDLE)
    cuadro_texto(slide, Inches(12.4), Inches(7.0), Inches(0.7), Inches(0.4),
                 str(numero), tamaño=12, color=COLORES["texto"],
                 alineacion=PP_ALIGN.RIGHT)


def tarjeta(slide, left, top, width, height, titulo, lineas):
    """Tarjeta gris con título en negrita y líneas de texto debajo."""
    rectangulo(slide, left, top, width, height, COLORES["claro"])
    rectangulo(slide, left, top, Inches(0.08), height, COLORES["acento"])
    caja = slide.shapes.add_textbox(left + Inches(0.3), top + Inches(0.2),
                                    width - Inches(0.5), height - Inches(0.4))
    tf = caja.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = titulo
    r.font.size = Pt(18)
    r.font.bold = True
    r.font.color.rgb = COLORES["primario"]
    r.font.name = FUENTE
    for linea in lineas:
        p = tf.add_paragraph()
        p.space_before = Pt(6)
        r = p.add_run()
        r.text = "• " + linea
        r.font.size = Pt(14)
        r.font.color.rgb = COLORES["texto"]
        r.font.name = FUENTE


# ---------------------------------------------------------------------------
# DIAPOSITIVAS
# ---------------------------------------------------------------------------


def diapositiva_portada(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # layout en blanco
    fondo(slide, COLORES["primario"])
    rectangulo(slide, 0, Inches(5.6), ANCHO, Inches(0.12), COLORES["acento"])
    cuadro_texto(slide, Inches(1), Inches(2.0), Inches(11.3), Inches(1.6),
                 TITULO, tamaño=48, negrita=True, color=COLORES["blanco"],
                 alineacion=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    cuadro_texto(slide, Inches(1), Inches(3.6), Inches(11.3), Inches(0.7),
                 SUBTITULO, tamaño=24, color=COLORES["acento"],
                 alineacion=PP_ALIGN.CENTER)
    cuadro_texto(slide, Inches(1), Inches(5.9), Inches(11.3), Inches(0.5),
                 AUTOR, tamaño=18, color=COLORES["blanco"],
                 alineacion=PP_ALIGN.CENTER)
    cuadro_texto(slide, Inches(1), Inches(6.4), Inches(11.3), Inches(0.5),
                 FECHA, tamaño=16, color=COLORES["blanco"],
                 alineacion=PP_ALIGN.CENTER)


def diapositiva_indice(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    encabezado(slide, "Índice", len(prs.slides._sldIdLst))
    elementos = [
        (0, "Introducción"),
        (0, "Objetivos"),
        (0, "Desarrollo"),
        (0, "Resultados"),
        (0, "Conclusiones"),
        (0, "Preguntas"),
    ]
    lista_viñetas(slide, Inches(2.5), Inches(1.8), Inches(8), Inches(4.5),
                  elementos, tamaño=24)


def diapositiva_introduccion(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    encabezado(slide, "Introducción", len(prs.slides._sldIdLst))
    texto = (
        "Presenta aquí el contexto y la motivación de tu trabajo: qué problema "
        "abordas, por qué es relevante y qué esperas lograr con esta "
        "presentación.\n\n"
        "Mantén los párrafos breves: 3 o 4 ideas por diapositiva bastan para "
        "que el mensaje sea claro y fácil de seguir."
    )
    caja = cuadro_texto(slide, Inches(1.2), Inches(2.0), Inches(10.9),
                        Inches(4.0), texto, tamaño=22)
    for p in caja.text_frame.paragraphs:
        p.space_after = Pt(14)


def diapositiva_objetivos(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    encabezado(slide, "Objetivos", len(prs.slides._sldIdLst))
    tarjeta(slide, Inches(0.9), Inches(1.9), Inches(5.7), Inches(4.6),
            "Objetivo general",
            ["Describe en una frase el logro principal que buscas.",
             "Debe ser claro, medible y acotado en el tiempo."])
    tarjeta(slide, Inches(6.9), Inches(1.9), Inches(5.7), Inches(4.6),
            "Objetivos específicos",
            ["Primer objetivo concreto y verificable.",
             "Segundo objetivo concreto y verificable.",
             "Tercer objetivo concreto y verificable."])


def diapositiva_desarrollo_1(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    encabezado(slide, "Desarrollo: Tema 1", len(prs.slides._sldIdLst))
    lista_viñetas(slide, Inches(1.2), Inches(1.9), Inches(11), Inches(4.8), [
        (0, "Primer punto principal del desarrollo."),
        (1, "Detalle o evidencia que respalda el punto."),
        (1, "Ejemplo concreto o dato de apoyo."),
        (0, "Segundo punto principal del desarrollo."),
        (1, "Detalle o evidencia que respalda el punto."),
        (0, "Tercer punto principal del desarrollo."),
    ], tamaño=22)


def diapositiva_desarrollo_2(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    encabezado(slide, "Desarrollo: Tema 2", len(prs.slides._sldIdLst))
    tarjeta(slide, Inches(0.9), Inches(1.9), Inches(5.7), Inches(4.6),
            "Enfoque / método",
            ["Cómo abordaste el tema.",
             "Herramientas o fuentes utilizadas.",
             "Pasos principales del proceso."])
    tarjeta(slide, Inches(6.9), Inches(1.9), Inches(5.7), Inches(4.6),
            "Hallazgos / ideas clave",
            ["Resultado más importante obtenido.",
             "Segunda idea relevante.",
             "Tercera idea relevante."])


def diapositiva_resultados(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    encabezado(slide, "Resultados", len(prs.slides._sldIdLst))
    datos = CategoryChartData()
    datos.categories = ["Trimestre 1", "Trimestre 2", "Trimestre 3", "Trimestre 4"]
    datos.add_series("Serie A", (4.5, 6.0, 7.5, 9.0))
    datos.add_series("Serie B", (3.0, 4.2, 5.5, 7.0))
    grafico = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(1.2), Inches(1.8), Inches(7.5), Inches(4.8), datos,
    ).chart
    grafico.has_legend = True
    grafico.legend.position = XL_LEGEND_POSITION.BOTTOM
    grafico.legend.include_in_layout = False
    for serie, color in zip(grafico.series, (COLORES["primario"], COLORES["acento"])):
        serie.format.fill.solid()
        serie.format.fill.fore_color.rgb = color
    tarjeta(slide, Inches(9.1), Inches(2.2), Inches(3.5), Inches(4.0),
            "Lectura rápida",
            ["Serie A crece de forma sostenida.",
             "Serie B acompaña con menor magnitud.",
             "Reemplaza estos datos por los tuyos."])


def diapositiva_tabla(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    encabezado(slide, "Datos clave", len(prs.slides._sldIdLst))
    filas, columnas = 5, 3
    tabla = slide.shapes.add_table(
        filas, columnas, Inches(1.5), Inches(2.0), Inches(10.3), Inches(3.6)
    ).table
    cabeceras = ["Categoría", "Descripción", "Valor"]
    cuerpo = [
        ["Aspecto 1", "Descripción breve del aspecto.", "00 %"],
        ["Aspecto 2", "Descripción breve del aspecto.", "00 %"],
        ["Aspecto 3", "Descripción breve del aspecto.", "00 %"],
        ["Aspecto 4", "Descripción breve del aspecto.", "00 %"],
    ]
    for j, texto in enumerate(cabeceras):
        celda = tabla.cell(0, j)
        celda.text = texto
        celda.fill.solid()
        celda.fill.fore_color.rgb = COLORES["primario"]
        r = celda.text_frame.paragraphs[0].runs[0]
        r.font.bold = True
        r.font.size = Pt(16)
        r.font.color.rgb = COLORES["blanco"]
        r.font.name = FUENTE
        if j == 2:
            celda.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    for i, fila in enumerate(cuerpo, start=1):
        for j, texto in enumerate(fila):
            celda = tabla.cell(i, j)
            celda.text = texto
            celda.fill.solid()
            celda.fill.fore_color.rgb = (
                COLORES["claro"] if i % 2 else COLORES["blanco"]
            )
            r = celda.text_frame.paragraphs[0].runs[0]
            r.font.size = Pt(14)
            r.font.color.rgb = COLORES["texto"]
            r.font.name = FUENTE
            if j == 2:  # columna "Valor" centrada
                celda.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    for j, ancho in enumerate((Inches(3.0), Inches(5.3), Inches(2.0))):
        tabla.columns[j].width = ancho
    cuadro_texto(slide, Inches(1.5), Inches(5.8), Inches(10.3), Inches(0.5),
                 "Nota: sustituye los datos de ejemplo por los resultados reales.",
                 tamaño=14, color=COLORES["texto"], alineacion=PP_ALIGN.CENTER)


def diapositiva_conclusiones(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    encabezado(slide, "Conclusiones", len(prs.slides._sldIdLst))
    lista_viñetas(slide, Inches(1.2), Inches(1.9), Inches(11), Inches(4.8), [
        (0, "Primera conclusión: resume el hallazgo más importante."),
        (0, "Segunda conclusión: conecta los resultados con los objetivos."),
        (0, "Tercera conclusión: indica el siguiente paso o recomendación."),
    ], tamaño=22)


def diapositiva_cierre(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fondo(slide, COLORES["primario"])
    rectangulo(slide, 0, Inches(4.6), ANCHO, Inches(0.12), COLORES["acento"])
    cuadro_texto(slide, Inches(1), Inches(2.4), Inches(11.3), Inches(1.4),
                 "¡Gracias!", tamaño=54, negrita=True, color=COLORES["blanco"],
                 alineacion=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    cuadro_texto(slide, Inches(1), Inches(3.7), Inches(11.3), Inches(0.7),
                 "¿Preguntas o comentarios?", tamaño=24, color=COLORES["acento"],
                 alineacion=PP_ALIGN.CENTER)
    cuadro_texto(slide, Inches(1), Inches(5.0), Inches(11.3), Inches(0.5),
                 CONTACTO, tamaño=16, color=COLORES["blanco"],
                 alineacion=PP_ALIGN.CENTER)


# ---------------------------------------------------------------------------
# PROGRAMA PRINCIPAL
# ---------------------------------------------------------------------------


def main():
    salida = sys.argv[1] if len(sys.argv) > 1 else "presentacion.pptx"

    prs = Presentation()
    prs.slide_width, prs.slide_height = ANCHO, ALTO

    diapositiva_portada(prs)
    diapositiva_indice(prs)
    diapositiva_introduccion(prs)
    diapositiva_objetivos(prs)
    diapositiva_desarrollo_1(prs)
    diapositiva_desarrollo_2(prs)
    diapositiva_resultados(prs)
    diapositiva_tabla(prs)
    diapositiva_conclusiones(prs)
    diapositiva_cierre(prs)

    prs.save(salida)
    print(f"Presentación generada: {salida} ({len(prs.slides._sldIdLst)} diapositivas)")


if __name__ == "__main__":
    main()
